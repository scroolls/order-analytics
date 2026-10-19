"""Export analysis results to Excel, CSV and JSON."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill


SUMMARY_LABELS = {
    "orders": "Всего заказов",
    "completed_orders": "Завершённые заказы",
    "cancelled_orders": "Отменённые заказы",
    "returned_orders": "Возвраты",
    "net_revenue": "Выручка после возвратов, ₽",
    "return_rate": "Доля возвратов",
    "return_rate_ci_low": "Нижняя граница 95% интервала",
    "return_rate_ci_high": "Верхняя граница 95% интервала",
    "late_rate": "Доля доставок позже срока",
    "median_delivery_days": "Медиана срока доставки, дни",
    "p90_delivery_days": "90-й процентиль срока доставки, дни",
}
SHEETS = {
    "monthly": "Динамика",
    "categories": "Категории",
    "regions": "Регионы",
    "channels": "Каналы",
}
COLUMN_LABELS = {
    "month": "Месяц",
    "category": "Категория",
    "region": "Регион",
    "channel": "Канал",
    "orders": "Заказы",
    "completed_orders": "Завершено",
    "returned_orders": "Возвраты",
    "net_revenue": "Выручка, ₽",
    "return_rate": "Доля возвратов",
    "late_rate": "Доля просрочек",
}


def write_outputs(
    data: pd.DataFrame,
    metrics: dict[str, pd.DataFrame],
    result: dict[str, float | int],
    output_dir: Path,
) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    bi_dir = output_dir / "power_bi"
    bi_dir.mkdir(exist_ok=True)

    export_data = data.copy()
    for column in ["order_date", "promised_date", "delivered_date"]:
        export_data[column] = export_data[column].dt.strftime("%Y-%m-%d").fillna("")
    export_data.to_csv(bi_dir / "orders.csv", index=False, encoding="utf-8-sig")
    for name, frame in metrics.items():
        frame.to_csv(output_dir / f"{name}.csv", index=False, encoding="utf-8-sig")
    (output_dir / "summary.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    summary_frame = pd.DataFrame(
        {"Показатель": [SUMMARY_LABELS[key] for key in result], "Значение": list(result.values())}
    )
    with pd.ExcelWriter(output_dir / "report.xlsx", engine="openpyxl") as writer:
        summary_frame.to_excel(writer, sheet_name="Сводка", index=False)
        for name, sheet in SHEETS.items():
            metrics[name].rename(columns=COLUMN_LABELS).to_excel(writer, sheet_name=sheet, index=False)

        workbook = writer.book
        for sheet in workbook:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 28
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="17324D")
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(vertical="center")
            for column in sheet.columns:
                width = max(len(str(cell.value or "")) for cell in column)
                sheet.column_dimensions[column[0].column_letter].width = min(max(width + 3, 14), 42)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.font = Font(name="Calibri", size=11, color="1B2733")
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.00" if isinstance(cell.value, float) else "#,##0"

        overview = workbook["Сводка"]
        overview.column_dimensions["A"].width = 39
        overview.column_dimensions["B"].width = 22
        for row_number in [7, 8, 9, 10]:
            overview.cell(row_number, 2).number_format = "0.0%"
        overview["A14"] = "95% интервал показывает неопределённость оценки доли возвратов."
        overview["A14"].font = Font(name="Calibri", size=10, italic=True, color="526677")

        monthly = workbook["Динамика"]
        for row in monthly.iter_rows(min_row=2, min_col=6, max_col=7):
            for cell in row:
                cell.number_format = "0.0%"
        monthly.column_dimensions["I"].width = 15

        revenue_chart = LineChart()
        revenue_chart.title = "Выручка по месяцам"
        revenue_chart.y_axis.title = "₽"
        revenue_chart.x_axis.title = "Месяц заказа"
        revenue_chart.add_data(Reference(monthly, min_col=5, min_row=1, max_row=monthly.max_row), titles_from_data=True)
        revenue_chart.set_categories(Reference(monthly, min_col=1, min_row=2, max_row=monthly.max_row))
        revenue_chart.width = 18
        revenue_chart.height = 9
        monthly.add_chart(revenue_chart, "I2")

        return_chart = BarChart()
        return_chart.title = "Доля возвратов по месяцам"
        return_chart.y_axis.title = "Доля"
        return_chart.add_data(Reference(monthly, min_col=6, min_row=1, max_row=monthly.max_row), titles_from_data=True)
        return_chart.set_categories(Reference(monthly, min_col=1, min_row=2, max_row=monthly.max_row))
        return_chart.width = 18
        return_chart.height = 9
        monthly.add_chart(return_chart, "I20")

        for name in ["Категории", "Регионы", "Каналы"]:
            sheet = workbook[name]
            for cell in sheet[1]:
                if cell.value in {"Доля возвратов", "Доля просрочек"}:
                    for row in sheet.iter_rows(min_row=2, min_col=cell.column, max_col=cell.column):
                        row[0].number_format = "0.0%"
